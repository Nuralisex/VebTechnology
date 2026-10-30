from pathlib import Path
import json
from docx import Document

required = ["ЦЕЛЬ РАБОТЫ", "ИНСТРУМЕНТЫ И ТЕХНОЛОГИИ", "ХОД РАБОТЫ", "РЕЗУЛЬТАТ РАБОТЫ", "ЗАКЛЮЧЕНИЕ"]
reports = sorted(Path(".").glob("Отчет_ЛР*.docx"))
result = []
for path in reports:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    sizes = sorted({r.font.size.pt for p in doc.paragraphs for r in p.runs if r.text.strip() and r.font.size})
    fonts = sorted({r.font.name for p in doc.paragraphs for r in p.runs if r.text.strip() and r.font.name})
    result.append({
        "file": path.name,
        "paragraphs": len(doc.paragraphs),
        "missing_sections": [section for section in required if section not in text],
        "cover": doc.paragraphs[0].text,
        "title": doc.paragraphs[10].text,
        "topic": doc.paragraphs[11].text,
        "font_sizes": sizes,
        "fonts": fonts,
        "bytes": path.stat().st_size,
    })
print(json.dumps(result, ensure_ascii=True, indent=2))
